"""Local document extraction with source provenance and no persistence of content."""

from __future__ import annotations

import csv
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .domain import DocumentItem, ExtractionArtifact, ExtractionStatus


class ExtractionError(RuntimeError):
    """The selected document cannot be extracted locally."""


# Rasterization scale for PDF page images; higher improves fine-print (e.g. depth ticks) legibility at higher token cost.
PDF_RENDER_SCALE = 2.5


def extract_document(document: DocumentItem) -> ExtractionArtifact:
    """Read one supported document into a bounded, provenance-labelled artifact."""
    if not document.is_ready:
        raise ExtractionError(f"Cannot extract a document with status '{document.status.value}'.")

    extractors = {
        ".csv": _extract_csv,
        ".xlsx": _extract_xlsx,
        ".docx": _extract_docx,
        ".pdf": _extract_pdf,
        ".doc": _extract_legacy_doc,
    }
    try:
        return extractors[document.extension](document)
    except (OSError, PdfReadError, ValueError) as error:
        raise ExtractionError(f"Could not read '{document.path.name}': {error}") from error


def render_pdf_pages(document: DocumentItem) -> tuple[bytes, ...]:
    """Render every page of a ready PDF to in-memory PNG bytes."""
    if not document.is_ready:
        raise ExtractionError(f"Cannot render a document with status '{document.status.value}'.")
    if document.extension != ".pdf":
        raise ExtractionError("Image processing is available only for PDF documents.")

    try:
        with pymupdf.open(document.path) as source:
            pages = tuple(page.get_pixmap(matrix=pymupdf.Matrix(PDF_RENDER_SCALE, PDF_RENDER_SCALE), alpha=False).tobytes("png") for page in source)
    except (OSError, RuntimeError, ValueError) as error:
        raise ExtractionError(f"Could not render '{document.path.name}': {error}") from error

    if not pages:
        raise ExtractionError(f"Could not render '{document.path.name}': the PDF has no pages.")
    return pages


def _extract_csv(document: DocumentItem) -> ExtractionArtifact:
    rows: list[str] = []
    locations: list[str] = []
    with document.path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
        for row_number, row in enumerate(csv.reader(source), start=1):
            rows.append(", ".join("" if value is None else value for value in row))
            locations.append(f"row {row_number}")
    return ExtractionArtifact(document, "\n".join(rows), tuple(locations))


def _extract_xlsx(document: DocumentItem) -> ExtractionArtifact:
    workbook = load_workbook(document.path, read_only=True, data_only=True)
    sections: list[str] = []
    locations: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                rows.append(", ".join("" if value is None else str(value) for value in row))
                locations.append(f"sheet '{sheet.title}', row {row_number}")
            sections.append(f"[{sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return ExtractionArtifact(document, "\n\n".join(sections), tuple(locations))


def _extract_docx(document: DocumentItem) -> ExtractionArtifact:
    source = Document(document.path)
    sections: list[str] = []
    locations: list[str] = []
    for paragraph_number, paragraph in enumerate(source.paragraphs, start=1):
        if paragraph.text.strip():
            sections.append(paragraph.text)
            locations.append(f"paragraph {paragraph_number}")
    for table_number, table in enumerate(source.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            sections.append(" | ".join(cell.text for cell in row.cells))
            locations.append(f"table {table_number}, row {row_number}")
    return ExtractionArtifact(document, "\n".join(sections), tuple(locations))


def _extract_pdf(document: DocumentItem) -> ExtractionArtifact:
    reader = PdfReader(document.path)
    pages: list[str] = []
    locations: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
            locations.append(f"page {page_number}")
    warnings: tuple[str, ...] = ()
    status = ExtractionStatus.COMPLETE
    if len(reader.pages) and not pages:
        warnings = ("No extractable PDF text found; this may be a scanned document and requires locally configured OCR.",)
        status = ExtractionStatus.REVIEW_REQUIRED
    return ExtractionArtifact(document, "\n\n".join(pages), tuple(locations), warnings, status)


def _extract_legacy_doc(document: DocumentItem) -> ExtractionArtifact:
    return ExtractionArtifact(
        document,
        "",
        (),
        ("Legacy .doc extraction requires an approved local conversion path and is not configured.",),
        ExtractionStatus.UNAVAILABLE,
    )