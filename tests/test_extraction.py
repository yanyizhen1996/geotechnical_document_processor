from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from document_processor.domain import DocumentStatus, ExtractionStatus, inspect_document
from document_processor.extraction import ExtractionError, extract_document, render_pdf_pages


def test_extract_csv_includes_rows_and_provenance(tmp_path: Path) -> None:
    path = tmp_path / "results.csv"
    path.write_text("name,value\nsoil,12\n", encoding="utf-8")

    artifact = extract_document(inspect_document(path))

    assert artifact.status is ExtractionStatus.COMPLETE
    assert artifact.content == "name, value\nsoil, 12"
    assert artifact.source_locations == ("row 1", "row 2")


def test_extract_xlsx_includes_sheet_and_row_provenance(tmp_path: Path) -> None:
    path = tmp_path / "results.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Testing"
    worksheet.append(["LL", 45])
    workbook.save(path)

    artifact = extract_document(inspect_document(path))

    assert "[Testing]" in artifact.content
    assert "LL, 45" in artifact.content
    assert artifact.source_locations == ("sheet 'Testing', row 1",)


def test_extract_docx_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "report.docx"
    source = Document()
    source.add_paragraph("Test summary")
    table = source.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "PI"
    table.rows[0].cells[1].text = "20"
    source.save(path)

    artifact = extract_document(inspect_document(path))

    assert "Test summary" in artifact.content
    assert "PI | 20" in artifact.content
    assert artifact.source_locations == ("paragraph 1", "table 1, row 1")


def test_extract_blank_pdf_requires_review_for_local_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as target:
        writer.write(target)

    artifact = extract_document(inspect_document(path))

    assert artifact.status is ExtractionStatus.REVIEW_REQUIRED
    assert "OCR" in artifact.warnings[0]


def test_render_pdf_pages_returns_a_png_for_each_page(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as target:
        writer.write(target)

    pages = render_pdf_pages(inspect_document(path))

    assert len(pages) == 2
    assert all(page.startswith(b"\x89PNG\r\n\x1a\n") for page in pages)


def test_extract_legacy_doc_reports_configured_limitation(tmp_path: Path) -> None:
    path = tmp_path / "legacy.doc"
    path.touch()

    artifact = extract_document(inspect_document(path))

    assert artifact.status is ExtractionStatus.UNAVAILABLE
    assert artifact.is_usable is False
    assert inspect_document(path).status is DocumentStatus.READY


def test_extract_malformed_pdf_returns_normalized_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.pdf"
    path.write_text("not a PDF", encoding="utf-8")

    try:
        extract_document(inspect_document(path))
    except ExtractionError as error:
        assert "broken.pdf" in str(error)
    else:
        raise AssertionError("Malformed PDF should fail extraction")